"""يولّد data/sales-data.json من ورقة Target. الاستخدام: python tools/build_json.py Total_2026.xlsx
(احفظ ملف xlsb بصيغة xlsx من Excel أولًا). يحتاج: pip install openpyxl"""
import json, sys
from openpyxl import load_workbook

ws = load_workbook(sys.argv[1], data_only=True)['Target ']
cats = ['Min', 'Syn', 'DEO', 'Fluids']; yrs = {2024: 3, 2025: 9, 2026: 15}
mn = ['Jan','Feb','Mar','Apr','May','Jun','Jul','Aug','Sep','Oct','Nov','Dec']
num = lambda v: float(v) if isinstance(v, (int, float)) else None
rows, issues = [], []

for i, m in enumerate(mn):
    for y, c0 in yrs.items():
        for k, cat in enumerate(cats):
            raw = ws.cell(4 + 2*i, c0 + k).value
            if isinstance(raw, str): issues.append(f"{y} {m} {cat}: النص '{raw.strip()}' بدل رقم؛ عوملت كقيمة مفقودة")
            rows.append(dict(year=y, month=i+1, monthName=m, quarter=i//3+1, half=1 if i < 6 else 2,
                             category=cat, target=num(ws.cell(3 + 2*i, c0 + k).value), actual=num(raw)))

src = {y: {'annualTarget': {c: num(ws.cell(27, c0+k).value) for k, c in enumerate(cats)},
           'annualActual': {c: num(ws.cell(28, c0+k).value) for k, c in enumerate(cats)}} for y, c0 in yrs.items()}

for y in yrs:
    for c in cats:
        s = sum(r['target'] or 0 for r in rows if r['year'] == y and r['category'] == c); st = src[y]['annualTarget'][c]
        if st is not None and abs(s - st) > 0.011: issues.append(f"{y} {c}: مجموع المستهدف الشهري {s:.2f} يختلف عن السنوي في المصدر {st}")

for y in yrs:
    for m in range(1, 13):
        a = [r['actual'] for r in rows if r['year'] == y and r['month'] == m]
        if all(v in (None, 0) for v in a) and (y != 2025):
            issues.append(f"{y} شهر {m}: الفعلي {'فارغ' if all(v is None for v in a) else 'صفر'} لكل الفئات؛ غير مؤكد أنه مبيعات صفرية")

json.dump(dict(unit="غير محدد في المصدر", source="Total_2026 / ورقة Target", asOf="حدّث التاريخ", rows=rows,
               sourceAnnual=src, issues=issues), open('data/sales-data.json', 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
